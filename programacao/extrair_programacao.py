"""Converte a planilha "Status + Programação" em tabelas normalizadas (CSV).

Cada aba "Fam ..." é uma pilha de blocos, um por produto. Cada bloco tem:
  - linha de cabeçalho (Entrada CQ, Código, Produto, Lote PA, Lote PI, ..., testes)
  - linha "Padrão" (status do padrão de referência de cada teste)
  - linhas de lote (SA/granel com Lote PI e PA/acabado com Lote PA)
Cada teste ocupa 3 colunas: Status | Prog. (turno 1/2/3 ou "s") | Backup.

Uso:
    python -m programacao.extrair_programacao "Status + Programação V0.xlsx" dados_programacao
"""

import csv
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

IGNORAR_TESTES = {"OBSERVAÇÃO", ""}
STATUS_CONCLUIDO = {
    "APROVADO", "N/A", "FEITO NO PI - OK", "FEITO NA VALIDAÇÃO", "PEGAR DO COA",
}


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(" ", "minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _turno(v):
    """Coluna Prog./Backup: 1, 2, 3 (ou "2-") = turno; "s" = feito."""
    t = _txt(v).upper().rstrip("-").strip()
    if t in ("1", "2", "3"):
        return t, ""
    return "", "s" if t == "S" else t


def _eh_cabecalho(ws, r):
    return _txt(ws.cell(r, 4).value) == "Entrada CQ" or _txt(ws.cell(r, 5).value) == "Código"


def extrair_familias(wb):
    lotes, analises, padroes, alocacao = [], [], [], []
    for ws in wb.worksheets:
        familia = ws.title.strip()
        if not familia.startswith("Fam"):
            continue
        testes, bloco = None, 0
        for r in range(1, ws.max_row + 1):
            cel = lambda c: ws.cell(r, c).value  # noqa: E731
            if _eh_cabecalho(ws, r):
                bloco += 1
                testes = {}
                for c in range(13, ws.max_column + 1):
                    nome = _txt(ws.cell(r, c).value)
                    if nome.upper() not in IGNORAR_TESTES:
                        testes[c] = nome
                continue
            if testes is None:
                continue
            produto = _txt(cel(6))
            if produto == "Padrão":
                for c, nome in testes.items():
                    if _txt(cel(c)):
                        padroes.append(dict(familia=familia, bloco=bloco, teste=nome,
                                            status=_txt(cel(c)).upper()))
                continue
            if not produto or not (_txt(cel(7)) or _txt(cel(8))):
                continue
            lote_pa, lote_pi = _txt(cel(7)), _txt(cel(8))
            lote_id = lote_pa or lote_pi
            obs = _txt(ws.cell(r, max(testes) + 3).value) if testes else ""
            lotes.append(dict(
                familia=familia, bloco=bloco, linha=r, lote=lote_id,
                tipo="PA" if lote_pa else "SA", codigo=_txt(cel(5)), produto=produto,
                lote_pa=lote_pa, lote_pi=lote_pi, entrada_cq=_txt(cel(4)),
                dias_cq=_txt(cel(9)), data_lib_prevista=_txt(cel(10)),
                cluster_inicial=_txt(cel(11)), cluster_atual=_txt(cel(12)),
                analista_t1=_txt(cel(1)), analista_t2=_txt(cel(2)), analista_t3=_txt(cel(3)),
            ))
            for c, nome in testes.items():
                status = _txt(cel(c)).upper()
                if not status:
                    continue
                turno_prog, prog = _turno(cel(c + 1))
                turno_bk, bk = _turno(cel(c + 2))
                analises.append(dict(
                    familia=familia, lote=lote_id, tipo="PA" if lote_pa else "SA",
                    lote_pi=lote_pi, teste=nome, status=status,
                    concluido="s" if status in STATUS_CONCLUIDO else "",
                    turno_programado=turno_prog, prog_ok=prog,
                    turno_backup=turno_bk, backup_ok=bk,
                    comentario=(ws.cell(r, c).comment.text.strip() if ws.cell(r, c).comment else ""),
                ))
            for t, c in (("1", 1), ("2", 2), ("3", 3)):
                if _txt(cel(c)):
                    alocacao.append(dict(familia=familia, lote=lote_id, turno=t, analista=_txt(cel(c))))
    return lotes, analises, padroes, alocacao


def extrair_escala(wb):
    ws = wb["Lista Analista Fam"]
    escala = []
    for r in range(2, ws.max_row + 1):
        for turno, (ca, cf) in (("1", (1, 2)), ("2", (4, 5)), ("3", (7, 8))):
            nome = _txt(ws.cell(r, ca).value)
            if nome:
                escala.append(dict(turno=turno, analista=nome,
                                   alocacao=_txt(ws.cell(r, cf).value)))
    return escala


def extrair_equipamentos(wb):
    ws = wb["Equipamentos"]
    cab = [_txt(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    out = []
    for r in range(2, ws.max_row + 1):
        linha = {cab[c - 1]: _txt(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        if linha.get("Equipamentos"):
            out.append(linha)
    return out


def extrair_dissolucao(wb):
    """Programação da dissolução: listas por turno a partir da linha 'Código'."""
    ws = wb["Dissolução"]
    blocos = {"1": 2, "2": 9, "3": 16}  # coluna de 'Código' de cada turno
    cab = next(r for r in range(1, ws.max_row + 1) if _txt(ws.cell(r, 2).value) == "Código")
    out = []
    for turno, c0 in blocos.items():
        for ordem, r in enumerate(range(cab + 1, ws.max_row + 1), start=1):
            lote = _txt(ws.cell(r, c0 + 2).value)
            if not lote:
                continue
            out.append(dict(turno=turno, ordem=ordem, codigo=_txt(ws.cell(r, c0).value),
                            produto=_txt(ws.cell(r, c0 + 1).value), lote=lote,
                            cluster=_txt(ws.cell(r, c0 + 3).value),
                            obs=_txt(ws.cell(r, c0 + 4).value)))
    return out


def _salvar(pasta, nome, linhas):
    if not linhas:
        return
    campos = list(dict.fromkeys(k for l in linhas for k in l))
    with open(pasta / nome, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=campos, delimiter=";")
        w.writeheader()
        w.writerows(linhas)


def main(arquivo, pasta_saida):
    wb = openpyxl.load_workbook(arquivo, data_only=True)
    pasta = Path(pasta_saida)
    pasta.mkdir(parents=True, exist_ok=True)
    lotes, analises, padroes, alocacao = extrair_familias(wb)
    tabelas = {
        "lotes.csv": lotes,
        "analises.csv": analises,
        "padroes.csv": padroes,
        "alocacao_lote.csv": alocacao,
        "escala.csv": extrair_escala(wb),
        "equipamentos.csv": extrair_equipamentos(wb),
        "programacao_dissolucao.csv": extrair_dissolucao(wb),
    }
    for nome, linhas in tabelas.items():
        _salvar(pasta, nome, linhas)
        print(f"{nome:28s} {len(linhas):5d} linhas")
    pend = [a for a in analises if not a["concluido"]]
    print(f"\n{len(lotes)} lotes, {len(analises)} análises, {len(pend)} pendentes, "
          f"{sum(1 for a in pend if a['turno_programado'])} programadas para um turno")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
